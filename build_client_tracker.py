"""Builds the client-facing list of delivered sets: ..\\RC\\RC_Shipped_Sets.xlsx

Deliberately narrow. This one leaves the workshop behind — no scores, no
statuses, no novelty or compliance numbers, no file paths, no bench. Four
columns the client can read at a glance, plus a genre count.

Shares its sources with build_shipping_tracker.py: the week folders are the
record of what went out, and rc_genres.json carries the subject genre. Run
that script first if a delivered set is missing a genre; it names the gaps.
"""
import json
import os
import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# build_shipping_tracker.py is a script, not a module: importing it would
# re-run the whole build and rewrite RC_Shipping_Tracker.xlsx as a side
# effect. Only its definitions are wanted here, so execute the file down to
# the point where it starts doing work. One parser, two trackers.
_SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "build_shipping_tracker.py")
_MARKER = "# ------------------------------------------------------------------ sources"


class _Helpers:
    pass


T = _Helpers()
with open(_SRC, encoding="utf-8") as _f:
    _defs = _f.read().split(_MARKER)[0]
_ns = {"__file__": _SRC, "__name__": "_bst_defs"}
exec(compile(_defs, _SRC, "exec"), _ns)
for _k, _v in _ns.items():
    if not _k.startswith("__"):
        setattr(T, _k, _v)

OUT = os.path.join(T.SHIP, "RC_Shipped_Sets.xlsx")

# A naive split on the first period breaks three openers the corpus actually
# uses: "C. S. Holling, writing in 1973...", "Mr. Whitcombe...", and any
# sentence carrying a decimal ("R0 comes out near 1.7, comfortably above...").
# Python cannot express "not preceded by any of these" as a lookbehind (it
# requires fixed width), so the check runs in code against the token before
# the stop.
_ABBREV = {"mr", "mrs", "ms", "dr", "prof", "st", "rev", "fr", "sr", "jr",
           "vs", "etc", "e.g", "i.e", "cf", "approx", "gen", "col", "capt",
           "lt", "sgt", "no", "vol", "pp", "ca", "bce", "ce", "ad", "bc",
           "u.s", "u.k", "fig", "op", "ed", "trans"}
# a stop, any closing quote, then whitespace and the next sentence's opener
_CANDIDATE = re.compile(r"[.!?][\"'”’)\]]?(?=\s+[\"'“‘(]?[A-Z0-9])")
_TOKEN_BEFORE = re.compile(r"([A-Za-z.]+)$")


def first_sentence(passage, hard_cap=520):
    """The opening line, as the client would read it off the page."""
    text = re.sub(r"\s+", " ", passage).strip()
    end = None
    for m in _CANDIDATE.finditer(text):
        if text[m.start()] == ".":
            tok = _TOKEN_BEFORE.search(text[:m.start()])
            word = tok.group(1).lower() if tok else ""
            if word in _ABBREV or (len(word) == 1 and word.isalpha()):
                continue                     # "Dr." / "C." — keep reading
            if text[m.start() - 1].isdigit() and m.start() + 1 < len(text) \
                    and text[m.start() + 1].isdigit():
                continue                     # a decimal, not a full stop
        end = m.end()
        break
    out = text[:end] if end else text
    if len(out) > hard_cap:                  # a very long opener gets trimmed
        cut = out.rfind(" ", 0, hard_cap)
        out = out[:cut if cut > 0 else hard_cap].rstrip(" ,;:") + " ..."
    return out


def week_label(wk):
    """Week_3_08_08 -> 'Week 3 (08 Aug 2026)'."""
    m = re.match(r"^Week_(\d+)_(\d{2})_(\d{2})$", wk)
    if not m:
        return wk
    months = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
    return "Week %s (%s %s 2026)" % (m.group(1), m.group(2), months[int(m.group(3)) - 1])


# ------------------------------------------------------------------ gather
genre_map = {}
if os.path.exists(T.GENRES):
    genre_map = json.load(open(T.GENRES, encoding="utf-8")).get("genres", {})
genre_by_canon = {T.canon(k): v for k, v in genre_map.items()}

# rebuild the shipped -> export match so each delivered set gets its genre
pool = T.load_corpus(T.EXP, "exported_rc_sets")
for rc, rec in T.load_corpus(T.MANUAL, "manual_rc_sets/used").items():
    pool.setdefault(rc, rec)
key_index = {}
for rc, rec in pool.items():
    for v in rec["passages"].values():
        if v:
            key_index.setdefault(T.norm(v)[:300], rc)

delivered = {}
for wk in sorted(os.listdir(T.SHIP)):
    if not wk.startswith("Week_"):
        continue                              # backup folders are not deliveries
    for dirpath, _dirs, files in os.walk(os.path.join(T.SHIP, wk)):
        for fn in sorted(files):
            if not fn.lower().endswith((".txt", ".docx")) or fn.startswith("~$"):
                continue
            sent = os.path.splitext(fn)[0]
            psg = T.split_passage(T.read_any(os.path.join(dirpath, fn)))
            r = delivered.setdefault((wk, sent),
                                     {"week": wk, "sent_id": sent, "passage": psg})
            if len(psg) > len(r["passage"]):  # weeks 1-4 ship .txt and .docx
                r["passage"] = psg

rows, missing = [], []
for (wk, sent), r in sorted(delivered.items(),
                            key=lambda kv: (T.week_date(kv[0][0]), kv[0][1])):
    tgt = T.norm(r["passage"])
    rc = key_index.get(tgt[:300])
    if rc is None:                            # passage edited after export
        best = max(((T.sim(tgt, T.norm(v)), k) for k, rec in pool.items()
                    for v in rec["passages"].values() if v), default=(0.0, None))
        rc = best[1] if best[0] >= 0.55 else None
    genre = genre_by_canon.get(T.canon(rc), "") if rc else ""
    if not genre:
        missing.append(sent)
    rows.append({"Week": week_label(wk), "RC ID": sent, "Genre": genre,
                 "Opening line": first_sentence(r["passage"])})

# ------------------------------------------------------------------ write
HDR_FILL = PatternFill("solid", start_color="1F4E78")
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY = Font(name="Calibri", size=11)
BAND = PatternFill("solid", start_color="F2F7FB")

wb = Workbook()
ws = wb.active
ws.title = "Shipped Sets"
cols = ["Week", "RC ID", "Genre", "Opening line"]
ws.append(cols)
for r in rows:
    ws.append([r[c] for c in cols])
for i in range(1, len(cols) + 1):
    c = ws.cell(row=1, column=i)
    c.fill, c.font = HDR_FILL, HDR_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
prev, band = None, False
for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
    if row[0].value != prev:                  # alternate shading per week
        band, prev = not band, row[0].value
    for c in row:
        c.font = BODY
        c.alignment = Alignment(vertical="top", wrap_text=(c.column == 4))
        if band:
            c.fill = BAND
for col, w in zip("ABCD", [22, 24, 28, 110]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:D%d" % (len(rows) + 1)

# ------------------------------------------------------- genre count sheet
wg = wb.create_sheet("Genres")
counts = Counter(r["Genre"] or "(unclassified)" for r in rows)
wg.append(["Genre", "Sets", "Share"])
for g, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
    wg.append([g, n, round(100.0 * n / max(len(rows), 1), 1)])
wg.append(["Total", len(rows), 100.0])
for i in range(1, 4):
    c = wg.cell(row=1, column=i)
    c.fill, c.font = HDR_FILL, HDR_FONT
    c.alignment = Alignment(horizontal="left")
for row in wg.iter_rows(min_row=2, max_row=len(counts) + 2):
    for c in row:
        c.font = BODY
    row[2].number_format = '0.0"%"'
wg.cell(row=len(counts) + 2, column=1).font = Font(name="Calibri", bold=True, size=11)
wg.cell(row=len(counts) + 2, column=2).font = Font(name="Calibri", bold=True, size=11)
for col, w in zip("ABC", [30, 8, 9]):
    wg.column_dimensions[col].width = w
wg.freeze_panes = "A2"

wb.save(OUT)
print("Saved %s" % OUT)
print("  %d delivered sets across %d weeks, %d genres"
      % (len(rows), len({r["Week"] for r in rows}), len(counts)))
if missing:
    print("  WARNING no genre for: %s" % ", ".join(sorted(missing)))
