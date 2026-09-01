"""Render delivery .txt sets to the .docx the client actually receives.

The Word files in the RC delivery folders are a plain 1:1 rendering of the text
files — one paragraph per line, no bold, theme default font, 1.25in margins.
Verified against Week_4_15_08/Hard/word/RC_HARD_260815_01.docx: every non-empty
line maps to exactly one paragraph with identical text.

That matters because the .docx is the deliverable and the .txt is the source. Any
edit to a .txt leaves its .docx stale, and nothing in the pipeline notices — you
would ship the old questions. Run this after editing text files.

Usage
-----
    # one file
    python tools/txt_to_docx.py "<RC>/Week_4_15_08/Hard/Text/RC_HARD_260815_01.txt"

    # every .txt under a week (writes into the sibling word/ folder)
    python tools/txt_to_docx.py "<RC>/Week_4_15_08"

    # check for stale/missing .docx without writing anything
    python tools/txt_to_docx.py "<RC>/Week_4_15_08" --check

A .txt at <...>/Text/NAME.txt is written to <...>/word/NAME.docx.
"""

from __future__ import annotations

import argparse
import io
import os
import sys

try:
    import docx
    from docx.shared import Inches
except ImportError:  # pragma: no cover
    sys.exit("python-docx is required:  pip install python-docx")

PAGE_MARGIN_IN = 1.25


def docx_path_for(txt_path: str) -> str:
    """<...>/Text/NAME.txt -> <...>/word/NAME.docx (falls back to alongside)."""
    d, name = os.path.split(txt_path)
    stem = os.path.splitext(name)[0] + ".docx"
    parent, leaf = os.path.split(d)
    if leaf.lower() == "text":
        return os.path.join(parent, "word", stem)
    return os.path.join(d, stem)


def lines_of(txt_path: str) -> list[str]:
    text = io.open(txt_path, encoding="utf-8", errors="replace").read()
    return [ln.rstrip() for ln in text.splitlines()]


def write_docx(txt_path: str, out_path: str) -> None:
    doc = docx.Document()
    for section in doc.sections:
        section.left_margin = Inches(PAGE_MARGIN_IN)
        section.right_margin = Inches(PAGE_MARGIN_IN)
    for line in lines_of(txt_path):
        doc.add_paragraph(line)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)


def docx_text(path: str) -> list[str]:
    return [p.text for p in docx.Document(path).paragraphs if p.text.strip()]


def is_stale(txt_path: str, out_path: str) -> str | None:
    """Returns a reason string when the .docx needs rebuilding, else None."""
    if not os.path.exists(out_path):
        return "missing"
    if docx_text(out_path) != [l for l in lines_of(txt_path) if l.strip()]:
        return "content differs from .txt"
    if os.path.getmtime(out_path) < os.path.getmtime(txt_path):
        return "older than .txt"
    return None


def collect(target: str) -> list[str]:
    if os.path.isfile(target):
        return [target]
    found = []
    for root, _dirs, files in os.walk(target):
        for f in sorted(files):
            if f.lower().endswith(".txt"):
                found.append(os.path.join(root, f))
    return found


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("target", help="a .txt file, or a folder to walk")
    ap.add_argument("--check", action="store_true",
                    help="report stale/missing .docx without writing")
    args = ap.parse_args(argv)

    txts = collect(args.target)
    if not txts:
        print(f"no .txt files under {args.target}")
        return 1

    stale = written = ok = 0
    for t in txts:
        out = docx_path_for(t)
        reason = is_stale(t, out)
        if args.check:
            if reason:
                stale += 1
                print(f"  STALE  {os.path.relpath(t, args.target)}  ({reason})")
            else:
                ok += 1
            continue
        if reason:
            write_docx(t, out)
            written += 1
            print(f"  wrote  {os.path.relpath(out, args.target)}  ({reason})")
        else:
            ok += 1

    if args.check:
        print(f"\n{len(txts)} .txt files: {ok} current, {stale} stale/missing")
        return 1 if stale else 0
    print(f"\n{len(txts)} .txt files: {written} rebuilt, {ok} already current")
    return 0


if __name__ == "__main__":
    sys.exit(main())
